#!/usr/bin/env python3
"""Approximate renderer for visual QA when LibreOffice is unavailable."""

from pathlib import Path
import sys

sys.path.insert(0, "/tmp/rapidlink_pptx_lib")

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "Tshwane_RapidLink_Hackathon_Pitch.pptx"
OUT = ROOT / "preview"
W, H = 1280, 720


def color_from(fmt, default=None):
    try:
        if fmt.type is None:
            return default
        return "#" + str(fmt.fore_color.rgb)
    except Exception:
        return default


def font_file(bold=False):
    return "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def wrap(draw, text, font, max_width):
    lines = []
    for raw in text.split("\n"):
        words = raw.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            test = current + " " + word
            if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
                current = test
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def render_text(draw, shp, sx, sy):
    text = shp.text.strip()
    if not text:
        return
    tf = shp.text_frame
    p = tf.paragraphs[0]
    run = p.runs[0] if p.runs else None
    pt = 14
    bold = False
    color = "#FFFFFF"
    if run is not None:
        try: pt = run.font.size.pt if run.font.size else pt
        except Exception: pass
        try: bold = bool(run.font.bold)
        except Exception: pass
        try:
            if run.font.color.rgb:
                color = "#" + str(run.font.color.rgb)
        except Exception: pass
    font = ImageFont.truetype(font_file(bold), max(7, int(pt * 1.22)))
    x = int(shp.left * sx) + 2
    y = int(shp.top * sy) + 2
    w = max(2, int(shp.width * sx) - 4)
    h = max(2, int(shp.height * sy) - 4)
    lines = wrap(draw, text, font, w)
    bbox = draw.textbbox((0, 0), "Ag", font=font)
    lh = max(8, bbox[3] - bbox[1] + 3)
    total = len(lines) * lh
    try:
        anchor = str(tf.vertical_anchor)
        if "MIDDLE" in anchor:
            y += max(0, (h-total)//2)
        elif "BOTTOM" in anchor:
            y += max(0, h-total)
    except Exception:
        pass
    try: align = str(p.alignment)
    except Exception: align = "LEFT"
    for line in lines:
        tw = draw.textbbox((0, 0), line, font=font)[2]
        tx = x
        if "CENTER" in align:
            tx = x + max(0, (w-tw)//2)
        elif "RIGHT" in align:
            tx = x + max(0, w-tw)
        draw.text((tx, y), line, font=font, fill=color)
        y += lh


def render_slide(prs, slide, idx):
    sx = W / prs.slide_width
    sy = H / prs.slide_height
    bg = color_from(slide.background.fill, "#0A1F44")
    img = Image.new("RGB", (W, H), bg)
    draw = ImageDraw.Draw(img)
    for shp in slide.shapes:
        x = int(shp.left*sx); y = int(shp.top*sy)
        w = int(shp.width*sx); h = int(shp.height*sy)
        fill = color_from(shp.fill, None) if hasattr(shp, "fill") else None
        outline = None
        try: outline = "#" + str(shp.line.color.rgb) if shp.line.color.rgb else None
        except Exception: pass
        if shp.shape_type == MSO_SHAPE_TYPE.LINE:
            draw.line((x, y, x+w, y+h), fill=outline or "#7F91AF", width=max(1, int(shp.line.width.pt/2)))
        elif shp.shape_type == MSO_SHAPE_TYPE.FREEFORM:
            if fill:
                draw.rounded_rectangle((x,y,x+w,y+h), radius=28, fill=fill, outline=outline)
        elif shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            typ = shp.auto_shape_type
            box=(x,y,x+w,y+h)
            if typ == MSO_SHAPE.OVAL:
                draw.ellipse(box, fill=fill, outline=outline, width=2)
            elif typ == MSO_SHAPE.ISOSCELES_TRIANGLE:
                draw.polygon([(x+w//2,y),(x+w,y+h),(x,y+h)], fill=fill, outline=outline)
            elif typ == MSO_SHAPE.CHEVRON:
                draw.polygon([(x,y),(x+w*2//3,y),(x+w,y+h//2),(x+w*2//3,y+h),(x,y+h),(x+w//3,y+h//2)], fill=fill, outline=outline)
            elif typ == MSO_SHAPE.ROUNDED_RECTANGLE:
                draw.rounded_rectangle(box, radius=max(4,min(w,h)//7), fill=fill, outline=outline, width=2)
            else:
                draw.rectangle(box, fill=fill, outline=outline, width=2)
        if getattr(shp, "has_text_frame", False):
            render_text(draw, shp, sx, sy)
    path = OUT / f"slide-{idx}.png"
    img.save(path)
    return img


def main():
    OUT.mkdir(exist_ok=True)
    prs = Presentation(PPTX)
    slides = [render_slide(prs, s, i+1) for i,s in enumerate(prs.slides)]
    thumb_w, thumb_h = 640, 360
    sheet = Image.new("RGB", (thumb_w*2, thumb_h*4), "#05132D")
    for i,img in enumerate(slides):
        sheet.paste(img.resize((thumb_w,thumb_h), Image.Resampling.LANCZOS), ((i%2)*thumb_w,(i//2)*thumb_h))
    sheet.save(OUT / "contact-sheet.png")
    print(OUT / "contact-sheet.png")


if __name__ == "__main__":
    main()
