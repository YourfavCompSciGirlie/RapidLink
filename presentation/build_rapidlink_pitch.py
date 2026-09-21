#!/usr/bin/env python3
"""Build the Tshwane RapidLink hackathon pitch deck."""

from pathlib import Path
import sys

LIB_DIR = "/tmp/rapidlink_pptx_lib"
if LIB_DIR not in sys.path:
    sys.path.insert(0, LIB_DIR)

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
OUT_FILE = OUT_DIR / "Tshwane_RapidLink_Hackathon_Pitch.pptx"

NAVY = "0A1F44"
NAVY_2 = "102B58"
NAVY_3 = "173867"
RED = "C8102E"
RED_2 = "EE3651"
WHITE = "FFFFFF"
LIGHT = "F5F7FA"
MID = "B8C4D9"
MUTED = "7F91AF"
INK = "13213A"
GREEN = "42D3A7"
AMBER = "FFBE55"
FONT = "DejaVu Sans"


def rgb(value):
    return RGBColor.from_string(value)


def set_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def shape(slide, kind, x, y, w, h, fill, line=None, radius=False):
    if radius:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    return s


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.5, dash=False):
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def textbox(slide, text, x, y, w, h, size=20, color=WHITE, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font=FONT,
            margin=0, break_word=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = break_word
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rich_text(slide, runs, x, y, w, h, size=18, align=PP_ALIGN.LEFT,
              valign=MSO_ANCHOR.MIDDLE, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold in runs:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return box


def slide_title(slide, title, kicker=None, num=None):
    if kicker:
        textbox(slide, kicker.upper(), 0.68, 0.36, 4.5, 0.24, 9, RED_2, True)
    textbox(slide, title, 0.68, 0.62, 10.7, 0.55, 26, WHITE, True)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.68, 1.22, 0.72, 0.05, RED, None)
    if num:
        textbox(slide, f"{num:02d}", 12.05, 0.48, 0.6, 0.32, 10, MUTED, True, PP_ALIGN.RIGHT)


def footer(slide, num):
    textbox(slide, "TSHWANE RAPIDLINK", 0.68, 7.12, 2.2, 0.18, 7, MUTED, True)
    textbox(slide, f"{num} / 7", 12.0, 7.10, 0.65, 0.2, 7, MUTED, True, PP_ALIGN.RIGHT)


def icon_pin(slide, cx, cy, scale=1.0, color=RED, dot=WHITE):
    # A teardrop assembled from a circle and a small triangle.
    tail = shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, cx-0.14*scale, cy+0.05*scale,
                 0.28*scale, 0.30*scale, color, None)
    tail.rotation = 180
    shape(slide, MSO_SHAPE.OVAL, cx-0.18*scale, cy-0.13*scale,
          0.36*scale, 0.36*scale, color, None)
    shape(slide, MSO_SHAPE.OVAL, cx-0.055*scale, cy-0.005*scale,
          0.11*scale, 0.11*scale, dot, None)


def icon_clock(slide, cx, cy, r=0.26, color=WHITE, fill=None):
    shape(slide, MSO_SHAPE.OVAL, cx-r, cy-r, 2*r, 2*r, fill, color)
    line(slide, cx, cy, cx, cy-r*0.54, color, 2)
    line(slide, cx, cy, cx+r*0.42, cy+r*0.12, color, 2)
    shape(slide, MSO_SHAPE.OVAL, cx-0.025, cy-0.025, 0.05, 0.05, color, None)


def icon_phone(slide, x, y, w=0.48, h=0.78, color=WHITE, fill=NAVY_3):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, color, True)
    s.line.width = Pt(1.5)
    shape(slide, MSO_SHAPE.RECTANGLE, x+0.07, y+0.12, w-0.14, h-0.25, NAVY, None)
    shape(slide, MSO_SHAPE.OVAL, x+w/2-0.025, y+h-0.09, 0.05, 0.05, color, None)


def icon_signal(slide, x, y, color=WHITE):
    for i, ht in enumerate([0.16, 0.27, 0.39, 0.52]):
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x+i*0.13, y+0.55-ht,
              0.085, ht, color, None, True)


def icon_nodes(slide, x, y, color=WHITE):
    pts = [(x+0.10,y+0.16),(x+0.48,y+0.11),(x+0.28,y+0.48),(x+0.62,y+0.48)]
    for a,b in [(0,1),(0,2),(1,2),(1,3),(2,3)]:
        line(slide, pts[a][0], pts[a][1], pts[b][0], pts[b][1], color, 1.8)
    for px,py in pts:
        shape(slide, MSO_SHAPE.OVAL, px-0.06, py-0.06, 0.12, 0.12, color, None)


def icon_wifi_off(slide, x, y, color=WHITE):
    # Concentric arcs approximated with partial curves are not dependable in PPT,
    # so use radiating links plus a strike-through.
    shape(slide, MSO_SHAPE.OVAL, x+0.28, y+0.46, 0.11, 0.11, color, None)
    line(slide, x+0.34, y+0.43, x+0.16, y+0.25, color, 2)
    line(slide, x+0.34, y+0.43, x+0.51, y+0.25, color, 2)
    line(slide, x+0.16, y+0.25, x+0.03, y+0.12, color, 2)
    line(slide, x+0.51, y+0.25, x+0.64, y+0.12, color, 2)
    line(slide, x+0.03, y+0.04, x+0.65, y+0.59, RED_2, 3)


def icon_building(slide, x, y, color=WHITE):
    shape(slide, MSO_SHAPE.RECTANGLE, x+0.08, y+0.10, 0.52, 0.50, None, color)
    shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y-0.03, 0.68, 0.22, color, None)
    for i in range(3):
        shape(slide, MSO_SHAPE.RECTANGLE, x+0.15+i*0.14, y+0.21, 0.07, 0.29, color, None)
    line(slide, x+0.03, y+0.61, x+0.65, y+0.61, color, 2)


def icon_landmark(slide, x, y, color=WHITE):
    icon_pin(slide, x+0.34, y+0.24, 1.2, color, NAVY_2)
    line(slide, x+0.03, y+0.62, x+0.65, y+0.62, color, 1.6)


def icon_report(slide, x, y, color=WHITE):
    icon_phone(slide, x+0.12, y, 0.40, 0.62, color, NAVY_3)
    shape(slide, MSO_SHAPE.OVAL, x+0.40, y-0.05, 0.24, 0.24, RED, None)
    textbox(slide, "!", x+0.40, y-0.05, 0.24, 0.24, 12, WHITE, True, PP_ALIGN.CENTER)


def icon_ai(slide, x, y, color=WHITE):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x+0.06, y+0.07, 0.56, 0.42, None, color, True)
    for px in [x+0.20, x+0.46]:
        shape(slide, MSO_SHAPE.OVAL, px, y+0.21, 0.07, 0.07, color, None)
    line(slide, x+0.25, y+0.38, x+0.44, y+0.38, color, 1.6)
    line(slide, x+0.34, y-0.02, x+0.34, y+0.07, color, 1.6)
    shape(slide, MSO_SHAPE.OVAL, x+0.30, y-0.08, 0.08, 0.08, RED_2, None)


def icon_dispatch(slide, x, y, color=WHITE):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x+0.05, y+0.16, 0.54, 0.33, None, color, True)
    shape(slide, MSO_SHAPE.RECTANGLE, x+0.11, y+0.08, 0.22, 0.12, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x+0.17, y+0.02, 0.10, 0.12, RED, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x+0.12, y+0.055, 0.20, 0.05, RED, None)
    for cx in [x+0.18, x+0.48]:
        shape(slide, MSO_SHAPE.OVAL, cx-0.07, y+0.43, 0.14, 0.14, color, None)


def icon_track(slide, x, y, color=WHITE):
    line(slide, x+0.06, y+0.47, x+0.22, y+0.31, color, 2)
    line(slide, x+0.22, y+0.31, x+0.36, y+0.40, color, 2)
    line(slide, x+0.36, y+0.40, x+0.60, y+0.13, color, 2)
    for px,py in [(x+0.06,y+0.47),(x+0.22,y+0.31),(x+0.36,y+0.40),(x+0.60,y+0.13)]:
        shape(slide, MSO_SHAPE.OVAL, px-0.045, py-0.045, 0.09, 0.09, RED_2, None)


def bullet_row(slide, y, headline, body, number=None, x=0.80, w=6.0,
               card_fill=NAVY_2, compact=False):
    h = 0.88 if compact else 1.02
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, card_fill, None, True)
    if number is not None:
        shape(slide, MSO_SHAPE.OVAL, x+0.20, y+0.24, 0.42, 0.42, RED, None)
        textbox(slide, str(number), x+0.20, y+0.24, 0.42, 0.42, 12, WHITE, True, PP_ALIGN.CENTER)
        tx = x+0.78
    else:
        shape(slide, MSO_SHAPE.OVAL, x+0.23, y+0.37, 0.12, 0.12, RED_2, None)
        tx = x+0.52
    textbox(slide, headline, tx, y+0.14, w-(tx-x)-0.2, 0.28, 15, WHITE, True)
    textbox(slide, body, tx, y+0.42, w-(tx-x)-0.25, h-0.48, 10.5, MID, False,
            valign=MSO_ANCHOR.TOP)


def add_map_network(slide):
    # Abstract Tshwane-shaped region with a subtle road network.
    builder = slide.shapes.build_freeform(0.25, 0.35, scale=Inches(1))
    builder.add_line_segments([
        (1.15, 0.10), (3.20, 0.00), (4.75, 0.72), (5.15, 2.10),
        (4.58, 3.02), (5.02, 4.52), (4.05, 6.15), (2.18, 6.45),
        (0.78, 5.76), (0.18, 4.44), (0.55, 3.12), (0.02, 1.80),
    ], close=True)
    poly = builder.convert_to_shape(Inches(7.2), Inches(0.2))
    poly.fill.solid(); poly.fill.fore_color.rgb = rgb(NAVY_2)
    poly.line.color.rgb = rgb(NAVY_3); poly.line.width = Pt(1)
    roads = [
        (7.45,1.05,12.40,2.55),(7.20,2.15,12.05,5.95),(7.80,5.95,12.35,3.55),
        (8.15,0.65,9.15,6.55),(10.35,0.35,9.85,6.85),(11.90,0.90,8.05,4.90),
        (7.55,3.45,12.55,3.15),(8.00,5.10,12.10,5.35),(9.05,1.15,12.60,4.90),
    ]
    for a,b,c,d in roads:
        line(slide,a,b,c,d,NAVY_3,1.2)
    # Orbital ring.
    ring = shape(slide, MSO_SHAPE.OVAL, 8.2, 1.05, 3.65, 4.85, None, NAVY_3)
    ring.line.width = Pt(1.2)
    for x,y,s in [(9.15,1.40,1.25),(11.45,2.38,1.0),(8.55,4.40,0.9),(10.55,5.35,1.15)]:
        shape(slide, MSO_SHAPE.OVAL, x-0.22*s, y-0.22*s, 0.44*s, 0.44*s, NAVY_3, None)
        icon_pin(slide,x,y,s,RED_2,WHITE)
    textbox(slide,"TSHWANE",9.15,3.25,2.2,0.35,10,MUTED,True,PP_ALIGN.CENTER)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Tshwane RapidLink"
    prs.core_properties.subject = "Tshwane Varsity Hackathon 2026 pitch"
    prs.core_properties.author = "Tshwane RapidLink Team"
    prs.core_properties.company = "Tshwane RapidLink"
    blank = prs.slide_layouts[6]

    # Slide 1 — Title
    slide = prs.slides.add_slide(blank); set_bg(slide)
    add_map_network(slide)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.72, 1.10, 0.12, 4.75, RED, None)
    textbox(slide, "TSHWANE", 1.12, 1.10, 5.7, 0.42, 18, RED_2, True)
    textbox(slide, "RapidLink", 1.08, 1.45, 6.1, 1.12, 45, WHITE, True)
    textbox(slide, "Every minute saved\nis a life saved.", 1.12, 2.68, 5.4, 1.30, 25, WHITE, True,
            valign=MSO_ANCHOR.TOP)
    textbox(slide, "Tshwane Varsity Hackathon 2026", 1.12, 4.28, 4.5, 0.35, 13, MID, False)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.12, 5.02, 3.65, 0.56, NAVY_2, NAVY_3, True)
    textbox(slide, "TEAM NAME  /  __________________", 1.36, 5.13, 3.15, 0.30, 10, MID, True)
    textbox(slide, "CONNECTED EMERGENCY RESPONSE", 1.12, 6.55, 4.2, 0.25, 8, MUTED, True)
    textbox(slide, "01", 12.05, 6.80, 0.6, 0.3, 10, MUTED, True, PP_ALIGN.RIGHT)

    # Slide 2 — Background
    slide = prs.slides.add_slide(blank); set_bg(slide)
    slide_title(slide, "Background", "The response gap", 2)
    # Left metric stack
    metrics = [
        ("3.5M+", "residents depend on one citywide emergency ecosystem"),
        ("SILOS", "police, fire, ambulance and communities lack one shared view"),
        ("5×", "slower response can separate township and affluent areas"),
    ]
    for i,(big,desc) in enumerate(metrics):
        y = 1.62+i*1.36
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.68, y, 4.35, 1.10, NAVY_2, None, True)
        textbox(slide,big,0.92,y+0.15,1.08,0.45,21,RED_2,True)
        textbox(slide,desc,2.05,y+0.17,2.65,0.67,11.5,WHITE,i==1,valign=MSO_ANCHOR.TOP)
    # Split-area visual card
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 5.35, 1.62, 7.30, 4.62, LIGHT, None, True)
    # township half
    shape(slide, MSO_SHAPE.RECTANGLE, 5.35, 1.62, 3.65, 4.62, "E8EDF4", None)
    textbox(slide,"TOWNSHIP",5.62,1.86,1.8,0.28,10,INK,True)
    textbox(slide,"40+ MIN",5.62,2.16,2.5,0.58,25,RED,True)
    # Road + informal homes
    line(slide,5.68,5.66,8.63,5.66,MUTED,2)
    for hx,hy,hw,hh in [(5.75,4.52,.72,.78),(6.52,4.16,.85,1.14),(7.44,4.64,.66,.66),(8.16,4.32,.55,.98)]:
        shape(slide,MSO_SHAPE.RECTANGLE,hx,hy,hw,hh,"B9C4D5",None)
        shape(slide,MSO_SHAPE.ISOSCELES_TRIANGLE,hx-0.05,hy-0.22,hw+0.10,0.30,MUTED,None)
    line(slide,6.12,3.05,6.12,4.38,MUTED,1.5)
    line(slide,5.95,3.28,6.55,3.28,MUTED,1.5)
    # affluent half
    textbox(slide,"AFFLUENT AREA",9.28,1.86,1.9,0.28,10,INK,True)
    textbox(slide,"8–12 MIN",9.28,2.16,2.6,0.58,25,NAVY,True)
    line(slide,9.32,5.66,12.32,5.66,MUTED,2)
    # trees/houses
    for hx in [9.46,10.56,11.56]:
        shape(slide,MSO_SHAPE.RECTANGLE,hx,4.50,.76,.82,WHITE,MUTED)
        shape(slide,MSO_SHAPE.ISOSCELES_TRIANGLE,hx-0.08,4.20,.92,.42,NAVY_3,None)
        shape(slide,MSO_SHAPE.RECTANGLE,hx+.29,4.89,.18,.43,NAVY_3,None)
    # clock divider
    shape(slide,MSO_SHAPE.OVAL,8.52,3.02,.96,.96,WHITE,None)
    icon_clock(slide,9.00,3.50,.34,RED,None)
    textbox(slide,"TIME = OUTCOME",7.80,6.43,2.4,.28,9,RED_2,True,PP_ALIGN.CENTER)
    footer(slide,2)

    # Slide 3 — Problem statement
    slide = prs.slides.add_slide(blank); set_bg(slide)
    slide_title(slide, "Problem Statement", "Fragmented in. Fragmented out.", 3)
    problems = [
        ("Too many entry points", "Calls, WhatsApp, social media and walk-ins."),
        ("Location is unclear", "Informal settlements often lack street addresses."),
        ("Dispatch is disconnected", "Each agency acts separately, with no coordination."),
        ("Access excludes people", "No airtime, data or address can mean no help."),
    ]
    for i,(head,body) in enumerate(problems):
        bullet_row(slide,1.52+i*1.14,head,body,i+1,0.68,6.05,NAVY_2,True)
    # Chaotic web diagram
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,7.08,1.52,5.58,4.84,NAVY_2,None,True)
    # incoming source nodes
    sources=[("CALL",7.42,2.00),("CHAT",9.05,1.75),("SOCIAL",10.80,2.05),("WALK-IN",11.15,4.70),("VOICE",9.10,5.28),("PHOTO",7.45,4.72)]
    center=(9.78,3.56)
    for label,x,y in sources:
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,1.12,.48,NAVY_3,None,True)
        textbox(slide,label,x,y,1.12,.48,9,WHITE,True,PP_ALIGN.CENTER)
        line(slide,x+.56,y+.24,center[0],center[1],RED_2,1.3,dash=True)
    # central tangle and phone
    for dx,dy in [(-.62,-.40),(.50,-.28),(-.52,.42),(.56,.46)]:
        shape(slide,MSO_SHAPE.OVAL,center[0]+dx-.08,center[1]+dy-.08,.16,.16,RED_2,None)
        line(slide,center[0]+dx,center[1]+dy,center[0],center[1],RED_2,1.3)
    shape(slide,MSO_SHAPE.OVAL,center[0]-.32,center[1]-.32,.64,.64,RED,None)
    textbox(slide,"?",center[0]-.32,center[1]-.32,.64,.64,25,WHITE,True,PP_ALIGN.CENTER)
    icon_phone(slide,11.35,2.86,.64,1.10,WHITE,NAVY)
    textbox(slide,"R0.00",11.42,3.18,.50,.22,8,RED_2,True,PP_ALIGN.CENTER)
    textbox(slide,"AIRTIME",11.33,3.43,.68,.18,6,WHITE,True,PP_ALIGN.CENTER)
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,7.64,6.08,4.46,.56,RED,None,True)
    textbox(slide,"THE SYSTEM BREAKS BEFORE HELP BEGINS",7.83,6.20,4.08,.27,10,WHITE,True,PP_ALIGN.CENTER)
    footer(slide,3)

    # Slide 4 — Solution flow
    slide = prs.slides.add_slide(blank); set_bg(slide)
    slide_title(slide, "The Solution", "One signal. One shared response.", 4)
    stages = [
        ("01","REPORT","Voice · text · photo\nUSSD · call",icon_report),
        ("02","TRIAGE","AI classification\nin under 5 seconds",icon_ai),
        ("03","LOCATE","Landmark match\nto GPS",icon_landmark),
        ("04","DISPATCH","Multi-agency\nin one click",icon_dispatch),
        ("05","TRACK","Updates · offline mode\nhospital pre-alert",icon_track),
    ]
    sx=0.68; card_w=2.30; gap=.19; y=1.66
    for i,(num,head,body,icon) in enumerate(stages):
        x=sx+i*(card_w+gap)
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,card_w,3.64,NAVY_2,None,True)
        textbox(slide,num,x+.20,y+.18,.40,.25,9,RED_2,True)
        shape(slide,MSO_SHAPE.OVAL,x+.78,y+.42,.74,.74,NAVY_3,None)
        icon(slide,x+.82,y+.49,WHITE)
        textbox(slide,head,x+.16,y+1.38,card_w-.32,.38,15,WHITE,True,PP_ALIGN.CENTER)
        textbox(slide,body,x+.18,y+1.90,card_w-.36,1.03,11,MID,False,PP_ALIGN.CENTER,MSO_ANCHOR.TOP)
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x+.40,y+3.18,card_w-.80,.05,RED,None,True)
        if i<4:
            # red chevron between cards
            chev=shape(slide,MSO_SHAPE.CHEVRON,x+card_w-.03,y+1.60,.25,.46,RED,None)
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,.98,5.72,11.36,.76,RED,None,True)
    rich_text(slide,[
        ("Works with zero airtime.  ",WHITE,True),
        ("No app needed.  ",WHITE,True),
        ("Keeps working when networks fail.",WHITE,True),
    ],1.18,5.88,10.96,.36,15,PP_ALIGN.CENTER)
    footer(slide,4)

    # Slide 5 — Competitors
    slide = prs.slides.add_slide(blank); set_bg(slide)
    slide_title(slide, "Existing Solutions", "Strong products. A different mandate.", 5)
    x0=.68; y0=1.56; widths=[2.15,4.35,5.12]; row_h=.74
    headers=["SOLUTION","WHAT IT DOES","LIMITATION"]
    xx=x0
    for j,w in enumerate(widths):
        shape(slide,MSO_SHAPE.RECTANGLE,xx,y0,w,.58,RED,None)
        textbox(slide,headers[j],xx+.18,y0+.10,w-.36,.32,10,WHITE,True)
        xx+=w+.03
    rows=[
        ("Namola","SOS app for private response","Needs smartphone + data"),
        ("AURA","Powers FNB, Investec, Tracker","Private subscription only"),
        ("Netstar","In-car panic button","Needs installed device"),
        ("Community Wolf","WhatsApp crime reporting","Not emergency dispatch"),
    ]
    for i,row in enumerate(rows):
        yy=y0+.61+i*(row_h+.05); xx=x0
        fill=NAVY_2 if i%2==0 else NAVY_3
        for j,(w,txt) in enumerate(zip(widths,row)):
            shape(slide,MSO_SHAPE.RECTANGLE,xx,yy,w,row_h,fill,None)
            textbox(slide,txt,xx+.18,yy+.10,w-.36,row_h-.18,12 if j==0 else 11,WHITE,j==0,
                    valign=MSO_ANCHOR.MIDDLE)
            xx+=w+.03
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,.68,5.65,11.72,.78,RED,None,True)
    textbox(slide,"ALL REQUIRE SMARTPHONE, DATA, CAR OR SUBSCRIPTION.",.98,5.77,6.30,.30,14,WHITE,True)
    textbox(slide,"NONE WORK WITH ZERO AIRTIME.",7.20,5.77,4.90,.30,14,WHITE,True,PP_ALIGN.RIGHT)
    footer(slide,5)

    # Slide 6 — Differentiators
    slide = prs.slides.add_slide(blank); set_bg(slide)
    slide_title(slide, "What Makes Us Different", "Access is the feature", 6)
    features=[
        ("01","NO-AIRTIME\nACCESS","Data → SMS → USSD → Call\nMTN *130*119# precedent",icon_signal),
        ("02","MULTI-AGENCY\nORCHESTRATION","The right mix of responders,\nnot just one",icon_nodes),
        ("03","LANDMARK\nLOCATION","Informal-settlement landmarks\nresolved to GPS",icon_landmark),
        ("04","OFFLINE-FIRST","Degrades gracefully when\nnetworks fail",icon_wifi_off),
        ("05","BUILT FOR\nTHE CITY","Public infrastructure,\nnot private subscription",icon_building),
    ]
    fw=2.30; fg=.19; fy=1.54
    for i,(num,head,body,icon) in enumerate(features):
        x=.68+i*(fw+fg)
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,fy,fw,4.80,NAVY_2,None,True)
        textbox(slide,num,x+.18,fy+.17,.44,.22,9,RED_2,True)
        shape(slide,MSO_SHAPE.OVAL,x+.76,fy+.53,.78,.78,RED,None)
        icon(slide,x+.81,fy+.62,WHITE)
        textbox(slide,head,x+.14,fy+1.57,fw-.28,.86,14,WHITE,True,PP_ALIGN.CENTER,
                MSO_ANCHOR.MIDDLE)
        shape(slide,MSO_SHAPE.RECTANGLE,x+.85,fy+2.54,.60,.04,RED,None)
        textbox(slide,body,x+.16,fy+2.83,fw-.32,1.12,10.2,MID,False,PP_ALIGN.CENTER,
                MSO_ANCHOR.TOP)
        textbox(slide,"CITY-SCALE",x+.30,fy+4.32,fw-.60,.22,8,RED_2,True,PP_ALIGN.CENTER)
    footer(slide,6)

    # Slide 7 — Commercialisation and ask
    slide = prs.slides.add_slide(blank); set_bg(slide)
    slide_title(slide, "How It Becomes Real", "Commercialisation & ask", 7)
    # Timeline rule
    line(slide,1.45,2.08,11.85,2.08,RED,3)
    phases=[
        ("PHASE 1","PILOT","One ward · 30 days","MTN zero-rated channel\nCity dispatch integration"),
        ("PHASE 2","CITY-WIDE","Tshwane licenses platform","MTN zero-rates permanently"),
        ("PHASE 3","NATIONAL","Johannesburg · Cape Town · Durban","Same model · repeatable rollout"),
    ]
    for i,(phase,title,lead,detail) in enumerate(phases):
        x=.74+i*4.14
        shape(slide,MSO_SHAPE.OVAL,x+1.55,1.89,.38,.38,RED,None)
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,2.38,3.73,1.50,NAVY_2,None,True)
        textbox(slide,phase,x+.24,2.55,1.0,.20,8,RED_2,True)
        textbox(slide,title,x+.24,2.80,3.25,.35,16,WHITE,True)
        textbox(slide,lead,x+.24,3.17,3.24,.25,10.5,WHITE,True)
        textbox(slide,detail,x+.24,3.43,3.24,.36,9.2,MID,False,valign=MSO_ANCHOR.TOP)
    textbox(slide,"WHO PAYS",.78,4.22,1.5,.26,10,RED_2,True)
    payers=[("CITY","Platform licensing"),("MTN","Zero-rating as social responsibility"),("CITIZEN","Never pays")]
    for i,(who,what) in enumerate(payers):
        x=.78+i*4.14
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,4.58,3.73,.76,NAVY_3,None,True)
        textbox(slide,who,x+.20,4.73,.90,.28,12,WHITE,True)
        textbox(slide,what,x+1.10,4.70,2.38,.34,10.2,MID,False,PP_ALIGN.RIGHT)
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,.78,5.75,11.82,.78,RED,None,True)
    textbox(slide,"ONE WARD. THIRTY DAYS.",1.10,5.90,4.2,.30,16,WHITE,True)
    textbox(slide,"Every minute saved is a life saved.",5.23,5.89,7.02,.32,17,WHITE,True,PP_ALIGN.RIGHT)
    footer(slide,7)

    prs.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    result = build_deck()
    print(result)
